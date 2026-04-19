"""Convert rendered resume HTML to a minimal ATS-friendly DOCX (no layout tables)."""

import re
from io import BytesIO
from typing import Optional

from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt


def _clean_text(text: Optional[str]) -> str:
    if not text:
        return ""
    return re.sub(r"\s+", " ", text).strip()


def html_to_docx(html: str) -> bytes:
    soup = BeautifulSoup(html, "html.parser")
    body = soup.body or soup
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def walk(node) -> None:
        name = getattr(node, "name", None)
        if name is None:
            return
        if name in ("script", "style", "svg", "noscript"):
            return
        if name in ("h1",):
            t = _clean_text(node.get_text())
            if t:
                doc.add_heading(t, level=0)
            return
        if name in ("h2",):
            t = _clean_text(node.get_text())
            if t:
                doc.add_heading(t, level=1)
            return
        if name in ("h3", "h4"):
            t = _clean_text(node.get_text())
            if t:
                doc.add_heading(t, level=2)
            return
        if name == "p":
            t = _clean_text(node.get_text())
            if t:
                doc.add_paragraph(t)
            return
        if name == "ul":
            for li in node.find_all("li", recursive=False):
                t = _clean_text(li.get_text())
                if t:
                    doc.add_paragraph(t, style="List Bullet")
            return
        if name == "ol":
            for li in node.find_all("li", recursive=False):
                t = _clean_text(li.get_text())
                if t:
                    doc.add_paragraph(t, style="List Number")
            return
        if name == "li":
            return
        if name == "br":
            return
        if name in ("div", "section", "article", "header", "footer", "main", "span", "body", "html"):
            for child in node.children:
                walk(child)
            return
        if name == "table":
            for tr in node.find_all("tr"):
                cells = [_clean_text(td.get_text()) for td in tr.find_all(["td", "th"])]
                line = " · ".join(c for c in cells if c)
                if line:
                    doc.add_paragraph(line)
            return

        for child in getattr(node, "children", []):
            walk(child)

    walk(body)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
